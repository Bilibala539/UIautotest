"""把所有 sheet 的状态分布列出来，写到 utf-8 文件"""
from openpyxl import load_workbook
from pathlib import Path
from collections import Counter

FILES = [
    r"C:\Users\wb.menglingzhao02\Desktop\首页其他功能_分sheet.xlsx",
    r"C:\Users\wb.menglingzhao02\Desktop\会话列表及会话盒子.xlsx",
]

out = []
for f in FILES:
    out.append("=" * 80)
    out.append(f"FILE: {f}")
    wb = load_workbook(f, data_only=True)
    out.append(f"All sheets: {wb.sheetnames}")
    for s in wb.sheetnames:
        ws = wb[s]
        # 最后一列通常是状态
        last_col = ws.max_column
        # 找表头
        header = [ws.cell(row=1, column=c).value for c in range(1, last_col + 1)]
        # 收集状态
        statuses = []
        for r in range(2, ws.max_row + 1):
            val = ws.cell(row=r, column=last_col).value
            if val is not None and str(val).strip():
                statuses.append(str(val).strip())
        cnt = Counter(statuses)
        out.append(f"\n  SHEET: {s}  (rows={ws.max_row}, status_col={last_col} = '{header[-1]}')")
        out.append(f"    used header: {header}")
        for status, n in cnt.most_common():
            short = status[:80] + ("..." if len(status) > 80 else "")
            out.append(f"    [{n}]  {short}")

Path(r"E:\UIAutotest\scripts\excel_sheets_summary.txt").write_text("\n".join(str(x) for x in out), encoding="utf-8")
print("done")
