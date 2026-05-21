"""列出两个 Excel 中每个 sheet 的表头和前 5 行内容，写到 utf-8 文件"""
from openpyxl import load_workbook
from pathlib import Path

FILES = [
    r"C:\Users\wb.menglingzhao02\Desktop\首页其他功能_分sheet.xlsx",
    r"C:\Users\wb.menglingzhao02\Desktop\会话列表及会话盒子.xlsx",
]

out = []
for f in FILES:
    out.append("=" * 80)
    out.append(f"FILE: {f}")
    wb = load_workbook(f, data_only=True)
    for s in wb.sheetnames:
        ws = wb[s]
        out.append(f"\n--- SHEET: {s}  (max_row={ws.max_row}, max_col={ws.max_column}) ---")
        for r in range(1, min(ws.max_row + 1, 6)):
            row = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            out.append(f"  row{r}: {row}")

Path(r"E:\UIAutotest\scripts\excel_structure.txt").write_text("\n".join(str(x) for x in out), encoding="utf-8")
print("written to scripts/excel_structure.txt")
