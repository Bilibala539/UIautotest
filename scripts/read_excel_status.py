"""读取Excel查看待实现的用例"""
import sys, os
sys.path.insert(0, r"E:\UIAutotest")
os.chdir(r"E:\UIAutotest")
from openpyxl import load_workbook

wb = load_workbook(r"C:\Users\wb.menglingzhao02\Desktop\会话列表及会话盒子.xlsx")
ws = wb.active

# Find all rows with status
print("Rows with status (checking columns):")
# First find column headers
headers = []
for col in range(1, ws.max_column + 1):
    val = ws.cell(row=1, column=col).value
    headers.append(val)
print(f"Headers: {headers}")

# Find the automation status column (last column or a specific one)
# Look for the column that contains our test file names
auto_col = None
for col in range(1, ws.max_column + 1):
    for row in range(2, min(10, ws.max_row + 1)):
        val = ws.cell(row=row, column=col).value
        if val and 'test_' in str(val):
            auto_col = col
            break
    if auto_col:
        break

if not auto_col:
    # Check last column
    auto_col = ws.max_column
    print(f"Using last column ({auto_col}) for automation status")

print(f"\nAutomation column: {auto_col}")
print(f"Total rows: {ws.max_row}")
print(f"\n--- All data rows ---")
for row in range(2, ws.max_row + 1):
    title = ws.cell(row=row, column=2).value or ws.cell(row=row, column=1).value or ""
    status = ws.cell(row=row, column=auto_col).value or ""
    if title:
        print(f"  Row {row:2d}: [{status[:30] if status else 'EMPTY':30s}] {title[:60]}")
