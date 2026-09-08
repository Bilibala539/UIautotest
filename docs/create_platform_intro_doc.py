from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(__file__).with_name("UIAutotest平台与Swan用例介绍.docx")


def set_cell_shading(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def set_cell_border(cell, color: str = "D9D9D9") -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)
        borders.append(element)
    properties.append(borders)


def configure_run(run, size: float = 10.5, bold: bool = False, color: str | None = None) -> None:
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_paragraph(document, text: str, style: str | None = None, bold_prefix: str | None = None) -> None:
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(6)
    if bold_prefix and text.startswith(bold_prefix):
        configure_run(paragraph.add_run(bold_prefix), bold=True)
        configure_run(paragraph.add_run(text[len(bold_prefix):]))
    else:
        configure_run(paragraph.add_run(text))


def add_bullet(document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    configure_run(paragraph.add_run(text))


def add_table(document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Cm(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "1F4E78")
        set_cell_border(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        configure_run(paragraph.add_run(header), size=9.5, bold=True, color="FFFFFF")
    for row_index, row in enumerate(rows):
        for index, value in enumerate(row):
            cell = table.add_row().cells[index]
            cell.width = Cm(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_shading(cell, "F4F8FB")
            set_cell_border(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            configure_run(paragraph.add_run(value), size=9.5)
    document.add_paragraph().paragraph_format.space_after = Pt(3)


def build_document() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    for style_name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12)):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    configure_run(title.add_run("UIAutotest 平台与 Swan 用例介绍"), size=22, bold=True)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    configure_run(subtitle.add_run("面向测试与研发同事的使用说明"), size=11, color="595959")
    subtitle.paragraph_format.space_after = Pt(18)

    document.add_heading("结论", level=1)
    add_paragraph(document, "框架应增加 Flutter 作为独立的平台入口，用于隔离 Swan 的测试用例、图片资产和运行配置；但不应重写一套底层驱动。Swan 的 Flutter 应用运行在 Android 设备上，继续复用 Android ADB、截图、OCR 和 AW 操作能力即可。")
    add_paragraph(document, "这样既能避免 TidyTree、智能音箱和 Swan 的资源混放，也能保证一次修复同时服务各个平台。")

    document.add_heading("平台与目录", level=1)
    add_table(
        document,
        ["运行平台", "对应项目", "测试用例", "正式图片", "图片草稿"],
        [
            ["web", "TidyTree", "tests/web/", "image/web/", "image_ai/web/"],
            ["android", "智能音箱", "tests/android/", "image/android/", "image_ai/android/"],
            ["flutter", "Swan", "tests/flutter/", "image/flutter/", "image_ai/flutter/"],
        ],
        [2.2, 2.5, 3.4, 3.7, 3.7],
    )
    add_paragraph(document, "新截取的模板图先放到 image_ai 对应平台目录；经过实际用例验证后，再使用 promote_image 转到 image 正式目录。图片不跨平台复用，以规避设备分辨率和视觉样式差异。")

    document.add_heading("AW 的复用边界", level=1)
    add_paragraph(document, "正式 pytest 用例只调用 user.aw，不直接写 ADB 命令、不固化点击坐标，也不重新实现 OCR。", bold_prefix="正式 pytest 用例")
    add_bullet(document, "固定文案：find_text_by_ocr 和 click_text")
    add_bullet(document, "Flutter 语义节点：accessibility_exists 和 click_accessibility")
    add_bullet(document, "通用交互：input_text、press_back、take_screenshot")
    add_paragraph(document, "AndroidAdbDriver 负责与设备通信并读取 UIAutomator 语义节点。FlutterAndroidDriver 继承它，只将平台名切换为 flutter，因此图片和测试配置会自动进入 Swan 的隔离目录。AWBase 仍是 OCR、模板匹配和点击的唯一实现。")

    document.add_heading("Swan 登录登出用例", level=1)
    add_paragraph(document, "用例文件为 tests/flutter/test_swan_login_logout.py。执行时会启动 Swan，判断当前是否已登录；如已登录则先退出，随后输入手机号和密码、勾选协议、登录并校验首页，最后退出并校验回到登录页。")
    add_paragraph(document, "诊断脚本 scripts/swan_login_logout_opencv_ocr.py 仍可用于录屏、UI 树转储和追踪日志，但文字定位、文字点击和语义节点点击已统一委托给 AWBase，不再维护独立 OCR 实现。")

    document.add_heading("如何运行", level=1)
    add_paragraph(document, "先连接 Android 设备并确认它显示为 device，然后在 PowerShell 中设置本次会话的环境变量：")
    command = document.add_paragraph()
    command.paragraph_format.left_indent = Cm(0.5)
    command.paragraph_format.space_after = Pt(8)
    configure_run(command.add_run('$env:PLATFORM = "flutter"\n$env:ANDROID_DEVICE_SERIAL = "设备序列号"\n$env:SWAN_PHONE = "测试手机号"\n$env:SWAN_PASSWORD = "测试密码"\nE:\\UIAutotest\\venv\\Scripts\\python.exe -m pytest tests/flutter/test_swan_login_logout.py -v'), size=9.5)
    add_paragraph(document, "账号信息只在本机环境变量或密钥管理系统中设置，不写入 .env 提交内容。未设置账号时，Swan 用例会跳过；ADB 不可用时，测试会明确报错。")

    document.add_heading("当前状态", level=1)
    add_bullet(document, "已完成：Flutter 平台入口、三类用例目录、三类图片目录、Swan pytest 用例和 AW 语义节点方法。")
    add_bullet(document, "已验证：Python 编译通过；OTP 服务测试通过，Swan 用例在非 Flutter 平台或未配置账号时会自动跳过。")
    add_bullet(document, "待实机回归：设备恢复连接后，配置授权测试账号即可完整执行登录和登出流程并保存运行证据。")

    document.save(OUTPUT)


if __name__ == "__main__":
    build_document()
