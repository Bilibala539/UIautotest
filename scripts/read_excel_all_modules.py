"""读取Excel - 查看所有用例功能点分布"""
import sys
sys.path.insert(0, r"E:\UIAutotest")
import openpyxl

wb = openpyxl.load_workbook(r"C:\Users\wb.menglingzhao02\Desktop\会话列表及会话盒子.xlsx", data_only=True)
ws = wb.active

# Collect all function points (module names)
modules = {}
for row in range(2, ws.max_row + 1):
    module = str(ws.cell(row=row, column=1).value or '').strip()
    title = str(ws.cell(row=row, column=3).value or '').strip()
    android = str(ws.cell(row=row, column=8).value or '').strip()
    
    if not title or title == 'None':
        continue
    
    if module not in modules:
        modules[module] = []
    modules[module].append({'row': row, 'title': title, 'android': android})

print("=== All modules and their cases ===\n")
for module, cases in modules.items():
    done = sum(1 for c in cases if c['android'] and c['android'] != 'None' and '跳过' not in c['android'])
    skipped = sum(1 for c in cases if '跳过' in c['android'])
    empty = sum(1 for c in cases if not c['android'] or c['android'] == 'None')
    print(f"[{module}] Total: {len(cases)} | Done: {done} | Skipped: {skipped} | Empty: {empty}")
    for c in cases:
        status_mark = "✅" if (c['android'] and c['android'] != 'None' and '跳过' not in c['android']) else "⏭️" if '跳过' in c['android'] else "❓"
        print(f"  {status_mark} Row {c['row']}: {c['title'][:60]}")
