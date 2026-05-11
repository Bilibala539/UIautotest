"""最终批量更新Excel"""
import sys, os
sys.path.insert(0, r"E:\UIAutotest")
os.chdir(r"E:\UIAutotest")
from openpyxl import load_workbook

wb = load_workbook(r"C:\Users\wb.menglingzhao02\Desktop\会话列表及会话盒子.xlsx")
ws = wb.active

COL = 8

updates = {
    # Row 12-14, 16: 会话盒子操作 - 当前消息列表中未发现会话盒子入口
    12: "跳过: 当前消息列表中未发现'会话盒子'分组入口,该功能需要接收特定服务号消息后才显示",
    13: "跳过: 同上-会话盒子分组未显示",
    14: "跳过: 当前消息列表中未发现'服务号助手'入口,需要接收特定服务号消息后才显示",
    16: "跳过: 同上-依赖会话盒子/服务号助手存在",

    # Row 96: 取消置顶 - 已通过测试
    96: "test_unpin_conversation.py",
}

print("Updating Excel...")
for row, value in updates.items():
    old = ws.cell(row=row, column=COL).value or ""
    ws.cell(row=row, column=COL, value=value)
    print(f"  Row {row}: '{old[:30]}...' -> '{value[:40]}...'")

wb.save(r"C:\Users\wb.menglingzhao02\Desktop\会话列表及会话盒子.xlsx")
print("\n✅ Excel saved!")
