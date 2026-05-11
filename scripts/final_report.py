"""生成最终统计报告"""
import sys, os
sys.path.insert(0, r"E:\UIAutotest")
os.chdir(r"E:\UIAutotest")
from openpyxl import load_workbook

wb = load_workbook(r"C:\Users\wb.menglingzhao02\Desktop\会话列表及会话盒子.xlsx")
ws = wb.active

COL = 8

# Categorize all rows
passed = []
skipped = []
empty = []

for row in range(2, ws.max_row + 1):
    test_point = ws.cell(row=row, column=3).value or ""
    status = ws.cell(row=row, column=COL).value or ""
    func = ws.cell(row=row, column=1).value or ""
    
    if not test_point:
        continue
    
    if "test_" in status:
        passed.append((row, func, test_point, status))
    elif "跳过" in status:
        skipped.append((row, func, test_point, status))
    elif status == "":
        empty.append((row, func, test_point))
    else:
        skipped.append((row, func, test_point, status))

print("=" * 80)
print("                     最终交付报告")
print("=" * 80)
print(f"\n📊 用例处理统计:")
print(f"  ✅ 通过 (生成测试文件): {len(passed)}")
print(f"  ⏭️ 跳过 (有原因说明): {len(skipped)}")
print(f"  ❓ 未处理: {len(empty)}")
print(f"  📋 总计: {len(passed) + len(skipped) + len(empty)}")

print(f"\n\n{'='*80}")
print("✅ 通过的用例:")
print(f"{'='*80}")
for row, func, tp, status in passed:
    print(f"  Row {row:3d}: [{status:40s}] {func[:20]} | {tp[:40]}")

print(f"\n\n{'='*80}")
print("⏭️ 跳过原因分类:")
print(f"{'='*80}")
# Group by reason type
reasons = {}
for row, func, tp, status in skipped:
    reason_type = status.split(":")[0] if ":" in status else status[:20]
    if reason_type not in reasons:
        reasons[reason_type] = []
    reasons[reason_type].append((row, tp[:40]))

for reason, items in sorted(reasons.items()):
    print(f"\n  📌 {reason} ({len(items)}条):")
    for row, tp in items[:5]:
        print(f"     Row {row}: {tp}")
    if len(items) > 5:
        print(f"     ... 还有 {len(items)-5} 条")

if empty:
    print(f"\n\n{'='*80}")
    print("❓ 未处理的用例:")
    print(f"{'='*80}")
    for row, func, tp in empty:
        print(f"  Row {row:3d}: {func[:20]} | {tp[:40]}")

print(f"\n\n{'='*80}")
print("📁 新增产出物:")
print(f"{'='*80}")
print(f"  测试用例文件:")
test_files = set(status for _, _, _, status in passed)
for f in sorted(test_files):
    print(f"    - tests/mumu/checklist/{f}")

print(f"\n  AW方法 (utils/aw/business/mumu.py):")
new_methods = [
    "enter_group_batch_manage()",
    "exit_group_batch_manage()",
    "select_group_in_batch(group_index)",
    "batch_select_all_groups()",
    "get_batch_selected_count()",
    "verify_batch_bottom_buttons()",
    "unpin_conversation(conversation_name)",
    "pin_conversation(conversation_name)",
]
for m in new_methods:
    print(f"    - {m}")

print(f"\n  PPUser方法 (utils/user.py):")
for m in new_methods:
    print(f"    - {m}")
