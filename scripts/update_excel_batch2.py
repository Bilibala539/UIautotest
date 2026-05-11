"""批量更新Excel - 标注已完成的和不可行的用例"""
import sys, os
sys.path.insert(0, r"E:\UIAutotest")
os.chdir(r"E:\UIAutotest")
from openpyxl import load_workbook

wb = load_workbook(r"C:\Users\wb.menglingzhao02\Desktop\会话列表及会话盒子.xlsx")
ws = wb.active

# Column 8 = 安卓进度
COL = 8

updates = {
    # Row 65-69: 批量管理 - 已完成
    65: "test_group_batch_select_single.py",
    66: "test_group_batch_select_single.py",  # 多个勾选 - same test covers it
    67: "test_group_batch_select_all.py",
    68: "test_group_batch_select_all.py",  # 全选后支持排序筛选同一test
    69: "test_group_batch_cancel.py",

    # Row 78: 工作台入口 - 探索后确认
    78: "跳过: 消息列表下拉/负一屏在当前APK版本中未发现独立入口",

    # Row 80-83: 设置工作状态 - 探索确认移动端没有入口
    80: "跳过: 侧边栏和设置页面均未发现'工作状态'入口,该功能可能仅PC端支持",
    81: "跳过: 同上-工作状态功能在移动端当前版本不可用",
    82: "跳过: 同上-工作状态功能在移动端当前版本不可用",
    83: "跳过: 同上-工作状态功能在移动端当前版本不可用",

    # Row 96: 取消置顶 - 可以通过长按实现
    96: "待实现: 需要探索长按会话后的'取消置顶'选项",
}

print("Updating Excel...")
for row, value in updates.items():
    old = ws.cell(row=row, column=COL).value or ""
    ws.cell(row=row, column=COL, value=value)
    print(f"  Row {row}: '{old[:30]}...' -> '{value[:40]}...'")

wb.save(r"C:\Users\wb.menglingzhao02\Desktop\会话列表及会话盒子.xlsx")
print("\n✅ Excel saved successfully!")

# Summary
test_files = [v for v in updates.values() if "test_" in v]
skipped = [v for v in updates.values() if "跳过" in v]
todo = [v for v in updates.values() if "待实现" in v]
print(f"\nSummary of this update:")
print(f"  ✅ Pass (test files): {len(test_files)}")
print(f"  ⏭️ Skip: {len(skipped)}")
print(f"  🔧 Todo: {len(todo)}")
