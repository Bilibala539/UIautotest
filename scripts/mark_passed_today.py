"""把今天早上回归 passed 的 row 也标到 Excel。"""
import re
from openpyxl import load_workbook

EXCEL = r"C:\Users\wb.menglingzhao02\Desktop\会话列表及会话盒子.xlsx"
SHEET = "已实现 (50)"

# 早上回归 passed 的行（从 progress log 提取）
# row8, 31, 32, 33, 34, 44, 45, 46, 49, 50, 51
PASSED_2026_05_21 = [8, 31, 32, 33, 34, 44, 45, 46, 49, 50, 51]

wb = load_workbook(EXCEL)
ws = wb[SHEET]

for r in PASSED_2026_05_21:
    v = ws.cell(r, 7).value or ""
    if not v.startswith("已通过") and not v.startswith("跳过"):
        ws.cell(r, 7).value = f"已通过 2026-05-21 - {v}"
        print(f'row{r}:', ws.cell(r, 7).value[:100])

wb.save(EXCEL)
print('done')
