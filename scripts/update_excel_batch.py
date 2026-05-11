"""批量更新Excel - 标记所有已通过的用例 + 标记新的不可执行用例"""
import sys, os
sys.path.insert(0, r"E:\UIAutotest")
os.chdir(r"E:\UIAutotest")
sys.stdout.reconfigure(encoding='utf-8')

import openpyxl
wb = openpyxl.load_workbook(r"C:\Users\wb.menglingzhao02\Desktop\会话列表及会话盒子.xlsx")
ws = wb.active

# Already passed tests
passed_cases = {
    71: "test_sidebar_display.py",
    58: "test_group_sort_default.py",
    59: "test_group_sort_recent.py",
    60: "test_group_sort_oldest.py",
    74: "test_sidebar_edit.py",
    90: "test_pin_contact.py",
}

for row, test_file in passed_cases.items():
    ws.cell(row=row, column=8).value = test_file

# Cases that need special environment or can't be automated based on exploration
skip_cases = {
    # Row 38-47: 会话添加到分组 - 需要PC端先配置分组
    38: "跳过: 会话分组功能需PC端预配置，当前环境群设置页面无'会话分组'入口",
    39: "跳过: 同上-会话分组功能需PC端预配置",
    40: "跳过: 同上-会话分组功能需PC端预配置",
    41: "跳过: 同上-会话分组功能需PC端预配置",
    42: "跳过: 同上-会话分组功能需PC端预配置",
    43: "跳过: 同上-会话分组功能需PC端预配置",
    44: "跳过: 同上-会话分组功能需PC端预配置",
    45: "跳过: 同上-会话分组功能需PC端预配置",
    46: "跳过: 同上-会话分组功能需PC端预配置",
    47: "跳过: 同上-会话分组功能需PC端预配置",
    # Row 62-63: 列表操作 - 退出/解散群组（破坏性操作，需要特定的可消耗群）
    62: "跳过: 退出群组为破坏性操作,会失去群组数据,不适合自动化反复执行",
    63: "跳过: 解散群组为破坏性操作,不可逆,不适合自动化",
    # Row 65-69: 批量管理 - 需要探索多选模态
    65: "待实现: 需要进一步探索群组批量管理的多选模态",
    66: "待实现: 需要进一步探索群组批量管理的多选模态",
    67: "待实现: 需要进一步探索群组批量管理的多选模态",
    68: "待实现: 需要进一步探索群组批量管理的多选模态",
    69: "待实现: 需要进一步探索群组批量管理的多选模态",
    # Row 72-73: 侧边栏 - 需要特定账号状态
    72: "跳过: 需要特定账号(用户名和签名超过一行+有徽章+有电子工卡+有侧边栏),当前账号不满足",
    73: "跳过: 侧边导航栏功能需要特定配置,当前账号未配置侧边导航栏应用",
    # Row 78: 工作台入口/负一屏
    78: "待实现: 需要探索消息列表下拉触发负一屏的方式",
    # Row 80-83: 工作状态
    80: "待实现: 需要找到工作状态设置入口(可能在侧边栏用户名区域)",
    81: "待实现: 同上-需先找到工作状态入口",
    82: "待实现: 同上-需先找到工作状态入口",
    83: "待实现: 同上-需先找到工作状态入口",
    # Row 92: 名称显示-需要名称超长的特定前置
    92: "跳过: 需要名称超长的特定会话前置条件,不容易构造",
    # Row 94: 排序-长按头像拖拽
    94: "跳过: 长按头像拖拽操作在模拟器上难以精确控制",
    # Row 96-97: 取消置顶
    96: "待实现: 需要验证取消置顶的具体操作路径",
    97: "跳过: 退群为破坏性操作,不适合自动化反复执行",
}

for row, status in skip_cases.items():
    current = ws.cell(row=row, column=8).value
    if current and ('test_' in str(current)):  # Don't overwrite passed tests
        continue
    ws.cell(row=row, column=8).value = status

wb.save(r"C:\Users\wb.menglingzhao02\Desktop\会话列表及会话盒子.xlsx")
print("Excel updated!")

# Print summary
print("\n=== 进度总结 ===")
total = 0
passed = 0
skipped = 0
todo = 0
for row_idx in range(2, 106):
    val = ws.cell(row=row_idx, column=8).value
    if ws.cell(row=row_idx, column=1).value is None:
        continue
    total += 1
    if val and 'test_' in str(val):
        passed += 1
    elif val and '跳过' in str(val):
        skipped += 1
    elif val and '待实现' in str(val):
        todo += 1

print(f"  总用例数: {total}")
print(f"  ✅ 已通过: {passed}")
print(f"  ⏭️ 跳过: {skipped}")
print(f"  🔧 待实现: {todo}")
print(f"  ❓ 未处理: {total - passed - skipped - todo}")
