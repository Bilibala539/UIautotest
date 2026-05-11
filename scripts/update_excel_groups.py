"""更新Excel - 标记会话分组用例为跳过，并开始处理可行的用例"""
import sys, os, logging, time
sys.path.insert(0, r"E:\UIAutotest")
os.chdir(r"E:\UIAutotest")
sys.stdout.reconfigure(encoding='utf-8')

import openpyxl
wb = openpyxl.load_workbook(r"C:\Users\wb.menglingzhao02\Desktop\会话列表及会话盒子.xlsx")
ws = wb.active

# Column H (index 8) = android status
COL_ANDROID = 8

# Update session group cases (rows 18-25) as skip
skip_reason_group = "跳过: 会话分组功能需要PC端预先配置分组才能在移动端显示，当前测试环境无PC端配置"
for row_idx in range(18, 26):
    current = ws.cell(row=row_idx, column=COL_ANDROID).value
    if current and '跳过' in str(current):
        continue  # Already marked
    ws.cell(row=row_idx, column=COL_ANDROID).value = skip_reason_group
    print(f"Row {row_idx}: Marked as skip (会话分组)")

# Update session group management cases (rows 27-36) as skip  
skip_reason_mgmt = "跳过: 分组管理功能需要PC端预先配置分组，当前测试环境消息列表页无分组Tab栏"
for row_idx in range(27, 37):
    current = ws.cell(row=row_idx, column=COL_ANDROID).value
    if current and '跳过' in str(current):
        continue
    ws.cell(row=row_idx, column=COL_ANDROID).value = skip_reason_mgmt
    print(f"Row {row_idx}: Marked as skip (分组管理)")

# Print remaining feasible cases
print("\n\n=== 可实现的用例 ===")
feasible_modules = ['会话盒子', '侧边栏tab', '设置工作状态', '置顶类型', '取消置顶', '名称显示', '排序']
for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
    row_data = list(row)
    if not row_data[0]:
        continue
    status = row_data[7] if len(row_data) > 7 else ''
    if status and '跳过' in str(status):
        continue
    module = row_data[0]
    name = row_data[2]
    steps = row_data[4]
    print(f"  Row {row_idx}: [{module}] {name}")
    print(f"         Steps: {steps}")

wb.save(r"C:\Users\wb.menglingzhao02\Desktop\会话列表及会话盒子.xlsx")
print("\n\nExcel updated and saved!")
