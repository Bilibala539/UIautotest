"""读取Excel用例的具体步骤描述"""
import sys, os, logging, time
sys.path.insert(0, r"E:\UIAutotest")
os.chdir(r"E:\UIAutotest")
sys.stdout.reconfigure(encoding='utf-8')

import openpyxl
wb = openpyxl.load_workbook(r"C:\Users\wb.menglingzhao02\Desktop\会话列表及会话盒子.xlsx")
ws = wb.active

# Print headers
headers = [cell.value for cell in ws[1]]
print("Headers:", headers)

# Print rows 18-36 (session group related cases)
print("\n=== 会话分组相关用例 (rows 18-36) ===")
for row in ws.iter_rows(min_row=18, max_row=36, values_only=True):
    # Print key columns
    row_data = list(row)
    if row_data[0]:  # Has case number
        print(f"\nRow: {row_data[0]}")
        print(f"  模块: {row_data[1]}")
        print(f"  用例名: {row_data[2]}")
        print(f"  前置: {row_data[3]}")
        print(f"  步骤: {row_data[4]}")
        print(f"  预期: {row_data[5]}")
