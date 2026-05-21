"""微调：补回 row2 的"已通过"状态，去掉 row3/4/5 重复的 '需第二台MuMu - '"""
import re
from openpyxl import load_workbook

EXCEL = r"C:\Users\wb.menglingzhao02\Desktop\会话列表及会话盒子.xlsx"
SHEET = "已实现 (50)"

wb = load_workbook(EXCEL)
ws = wb[SHEET]

# row2 补回"已通过"（昨天 2026-05-20 跑通的）
v = ws.cell(2, 7).value or ""
if not v.startswith("已通过"):
    ws.cell(2, 7).value = f"已通过 2026-05-20 - {v}"

# row3/4/5 去掉重复的 '需第二台MuMu - 需第二台MuMu - '
for r in (3, 4, 5):
    v = ws.cell(r, 7).value or ""
    v = re.sub(r'(需第二台MuMu - )(?=需第二台MuMu)', '', v)
    ws.cell(r, 7).value = v

wb.save(EXCEL)
print('done')

# 列出
import sys
sys.stdout.reconfigure(encoding='utf-8')
for r in (2, 3, 4, 5, 6, 7, 9, 42, 43):
    print(f'row{r}:', ws.cell(r, 7).value)
