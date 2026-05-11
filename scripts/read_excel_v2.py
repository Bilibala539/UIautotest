"""读取Excel - 更详细的内容"""
import sys, os
sys.path.insert(0, r"E:\UIAutotest")
os.chdir(r"E:\UIAutotest")
from openpyxl import load_workbook

wb = load_workbook(r"C:\Users\wb.menglingzhao02\Desktop\会话列表及会话盒子.xlsx")
ws = wb.active

# Print headers
print("=== HEADERS ===")
for col in range(1, 9):
    val = ws.cell(row=1, column=col).value
    print(f"  Col {col}: {val}")

# Print all data with col 3 (测试点) and col 8 (安卓进度)
print(f"\n=== DATA (col1=功能点, col3=测试点, col8=安卓进度) ===")
for row in range(2, min(106, ws.max_row + 1)):
    func = ws.cell(row=row, column=1).value or ""
    test_point = ws.cell(row=row, column=3).value or ""
    android_status = ws.cell(row=row, column=8).value or ""
    if test_point:
        flag = "✅" if "test_" in str(android_status) else ("⏭️" if android_status else "  ")
        print(f"  {flag} Row {row:3d}: [{android_status[:40] if android_status else '':40s}] {func[:20]:20s} | {test_point[:50]}")
