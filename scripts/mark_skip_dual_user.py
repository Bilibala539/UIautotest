"""把 row3/4/5（双用户）标记为跳过，避免回归脚本误跑"""
from openpyxl import load_workbook

EXCEL = r"C:\Users\wb.menglingzhao02\Desktop\会话列表及会话盒子.xlsx"
SHEET = "已实现 (50)"
SKIP_ROWS = (3, 4, 5)
PREFIX = "跳过 2026-05-21 - 需第二台MuMu - "

wb = load_workbook(EXCEL)
ws = wb[SHEET]
for r in SKIP_ROWS:
    v = ws.cell(r, 7).value
    if v and not v.startswith(PREFIX):
        ws.cell(r, 7).value = PREFIX + v
        print(f'row{r}: {v[:80]}')
wb.save(EXCEL)
print('done')
