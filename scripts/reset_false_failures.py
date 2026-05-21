"""撤销因 NODE_USERS 限制误判失败的行 —— 重置成原始的 '已实现 - ...' """
from openpyxl import load_workbook

EXCEL = r"C:\Users\wb.menglingzhao02\Desktop\会话列表及会话盒子.xlsx"
SHEET = "已实现 (50)"

wb = load_workbook(EXCEL)
ws = wb[SHEET]
for r in range(2, ws.max_row + 1):
    v = ws.cell(r, 7).value
    if not v:
        continue
    # 去掉 "失败 2026-05-20 - " 前缀
    if v.startswith("失败 2026-05-20 - "):
        new = v.replace("失败 2026-05-20 - ", "", 1)
        ws.cell(r, 7).value = new
        print(f'row{r} reset: {new[:80]}')
wb.save(EXCEL)
print('done')
