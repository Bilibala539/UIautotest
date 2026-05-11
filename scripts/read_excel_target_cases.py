"""读取Excel - 查看会话分组、工作状态、会话盒子相关用例"""
import sys
sys.path.insert(0, r"E:\UIAutotest")
import openpyxl

wb = openpyxl.load_workbook(r"C:\Users\wb.menglingzhao02\Desktop\会话列表及会话盒子.xlsx", data_only=True)
ws = wb.active

# Headers: {1: '功能点', 2: '用例级别', 3: '测试点/测试目的', 4: '前置条件', 5: '操作步骤', 6: '预期结果', 7: 'ui自动化ios进度', 8: '安卓进度'}
module_col = 1   # 功能点
title_col = 3    # 测试点/测试目的
precond_col = 4  # 前置条件
steps_col = 5    # 操作步骤
expected_col = 6 # 预期结果
android_col = 8  # 安卓进度

# Read all rows and filter by relevant modules
target_keywords = ['会话分组', '工作状态', '会话盒子', '分组管理', '分组栏', '隐藏栏']
all_cases = []

for row in range(2, ws.max_row + 1):
    module = str(ws.cell(row=row, column=module_col).value or '')
    title = str(ws.cell(row=row, column=title_col).value or '')
    steps = str(ws.cell(row=row, column=steps_col).value or '')
    expected = str(ws.cell(row=row, column=expected_col).value or '')
    precond = str(ws.cell(row=row, column=precond_col).value or '')
    android = str(ws.cell(row=row, column=android_col).value or '')
    
    if not title or title == 'None':
        continue
    
    # Check if this case belongs to our target modules
    is_target = False
    combined = module + title + steps + precond
    for kw in target_keywords:
        if kw in combined:
            is_target = True
            break
    
    if is_target:
        all_cases.append({
            'row': row,
            'module': module,
            'title': title,
            'steps': steps,
            'expected': expected,
            'precond': precond,
            'android': android
        })

print(f"Found {len(all_cases)} target cases:")
print("=" * 100)
for case in all_cases:
    status = case['android'] if case['android'] != 'None' else '(empty)'
    print(f"\nRow {case['row']} [{case['module']}]")
    print(f"  Title: {case['title'][:80]}")
    print(f"  Precond: {case['precond'][:80]}")
    print(f"  Steps: {case['steps'][:120]}")
    print(f"  Expected: {case['expected'][:120]}")
    print(f"  Android: {status}")